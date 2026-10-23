"""Сквозные проверки самостоятельных: время, изоляция, снимки и история."""
import io
import tempfile
from datetime import timedelta
from unittest.mock import patch
from PIL import Image
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from accounts.models import Profile
from problems.models import Task, TaskSubitem, TaskSource, Source, TaskImage
from .models import (Work, Position, Candidate, Run, Attempt, Answer,
                     AnswerCheckAttempt, Asset, Family)
from . import services as svc


@override_settings(INDEPENDENT_WORKS_ENABLED=True)
class IndependentWorkTests(TestCase):
    def setUp(self):
        self.media = tempfile.TemporaryDirectory()
        self.addCleanup(self.media.cleanup)
        settings = override_settings(MEDIA_ROOT=self.media.name)
        settings.enable()
        self.addCleanup(settings.disable)
        User = get_user_model()
        self.teacher = User.objects.create_user('teacher')
        Profile.objects.create(user=self.teacher, role='teacher', base_editor=True)
        self.student = User.objects.create_user('student')
        Profile.objects.create(user=self.student, teacher=self.teacher, role='student')
        self.other = User.objects.create_user('other')
        Profile.objects.create(user=self.other, role='student')
        self.work = Work.objects.create(owner=self.teacher, title='Уравнения')
        self.position = Position.objects.create(work=self.work)
        self.task = Task.objects.create(text_latex='Два плюс два?', answer_latex='SECRET_ANSWER', solution_latex='SECRET_SOLUTION', answer_mask='{{number: 4}}')
        self.candidate = Candidate.objects.create(position=self.position, task=self.task)

    def issued(self, **kwargs):
        run = svc.prepare(self.work, self.teacher, [self.student.pk], **kwargs)
        svc.issue(run.pk)
        return run.attempts.get()

    def begin(self, **kwargs):
        return svc.start(self.issued(**kwargs).pk)

    def test_unavailable_check_is_retried_limited_times_then_manual(self):
        import subprocess
        from assignments.checks import formulas
        self.task.answer_mask = '{{expression: a+28}}'
        self.task.save()
        a = self.begin()
        point = a.variant.points.get()
        formulas._checked.cache_clear()
        formulas._failed_inputs.clear()
        with patch.object(formulas, '_run_worker',
                          side_effect=subprocess.TimeoutExpired('worker', 3)):
            answer = svc.save_answer(a.pk, point.pk, 0, {'s1': '28+a'})
            self.assertEqual(answer.auto_check, '')
            self.assertEqual(answer.checks, {formulas.UNAVAILABLE_KEY: 1})
            svc.recheck_unavailable(a)
        answer.refresh_from_db()
        # Повторный сбой того же ввода — в ручную проверку, без вечных досверок
        self.assertEqual(answer.auto_check, 'unknown')
        self.assertEqual(answer.checks, {'s1': 'unknown', formulas.UNAVAILABLE_KEY: 2})
        with patch.object(formulas, '_run_worker') as worker:
            svc.recheck_unavailable(a)
            worker.assert_not_called()

    def test_cooldown_does_not_count_as_another_failed_check(self):
        import subprocess
        from assignments.checks import formulas
        self.task.answer_mask = '{{expression: a+163}}'
        self.task.save()
        attempt = self.begin()
        point = attempt.variant.points.get()
        formulas._checked.cache_clear()
        formulas._failed_inputs.clear()
        self.addCleanup(formulas._failed_inputs.clear)
        with patch.object(formulas, '_spawn',
                          side_effect=subprocess.TimeoutExpired('worker', 3)) as worker:
            answer = svc.save_answer(attempt.pk, point.pk, 0, {'s1': '163+a'})
            svc.recheck_unavailable(attempt)
            svc.recheck_unavailable(attempt)
            worker.assert_called_once()
        answer.refresh_from_db()
        self.assertEqual(answer.auto_check, '')
        self.assertEqual(answer.checks, {formulas.UNAVAILABLE_KEY: 1})
        formulas._failed_inputs.clear()
        svc.recheck_unavailable(attempt)
        answer.refresh_from_db()
        self.assertEqual(answer.auto_check, 'ok')
        self.assertEqual(answer.checks, {'s1': 'ok'})

    def url(self, name, *args):
        return reverse('independentworks:' + name, args=args)

    def test_unavailable_self_check_returns_saved_version_and_allows_next_save(self):
        from assignments.checks.formulas import FormulaBusy
        attempt = self.begin(self_check=True)
        point = attempt.variant.points.get()
        self.client.force_login(self.student)
        with patch.object(svc, 'check_mask', side_effect=FormulaBusy('занято')):
            response = self.client.post(self.url('check_answer', attempt.pk, point.pk), {
                'version': 0, f'p{point.pk}_s1': '4'})
        self.assertEqual(response.status_code, 200)
        result = response.json()
        self.assertFalse(result['ok'])
        self.assertEqual(result['version'], 1)
        self.assertFalse(AnswerCheckAttempt.objects.filter(answer__attempt=attempt).exists())
        response = self.client.post(self.url('save', attempt.pk, point.pk), {
            'version': result['version'], f'p{point.pk}_s1': '3'})
        self.assertEqual(response.status_code, 200)
        answer = attempt.answers.get()
        self.assertEqual((answer.version, answer.values), (2, {'s1': '3'}))

    def test_same_input_cannot_reset_failure_limit_by_saving_or_checking(self):
        from assignments.checks import formulas
        self.task.answer_mask = '{{number: 4}} {{expression: a+191}} {{number: 7}}'
        self.task.save()
        formulas._checked.cache_clear()
        attempt = self.begin(self_check=True)
        point = attempt.variant.points.get()
        values = {'s1': '4', 's2': '191+a', 's3': '8'}
        with patch.object(formulas, '_run_worker', side_effect=formulas.FormulaUnavailable('сбой')) as worker:
            answer = svc.save_answer(attempt.pk, point.pk, 0, values)
            answer = svc.save_answer(attempt.pk, point.pk, answer.version, values)
            self.assertEqual(answer.auto_check, 'unknown')
            self.assertEqual(answer.checks, {'s1': 'ok', 's2': 'unknown', 's3': 'wrong', formulas.UNAVAILABLE_KEY: 2})
            answer = svc.save_answer(attempt.pk, point.pk, answer.version, values)
            self.client.force_login(self.student)
            response = self.client.post(self.url('check_answer', attempt.pk, point.pk), {
                'version': answer.version, **{f'p{point.pk}_{k}': v for k, v in values.items()}})
            self.assertTrue(response.json()['ok'])
            self.assertEqual(worker.call_count, 2)
            # Изменение ввода сбрасывает счётчик и действительно запускает проверку.
            answer.refresh_from_db()
            values['s2'] = 'a+192'
            answer = svc.save_answer(attempt.pk, point.pk, answer.version, values)
            self.assertEqual(worker.call_count, 3)
            self.assertEqual(answer.checks, {formulas.UNAVAILABLE_KEY: 1})

    def test_teacher_pages_and_export_never_retry_formula_checks(self):
        attempt = self.begin()
        point = attempt.variant.points.get()
        Answer.objects.create(attempt=attempt, point=point, values={'s1': '4'},
                              version=1, auto_check='')
        self.client.force_login(self.teacher)
        with patch.object(svc, 'check_mask') as check:
            for name, args in [('progress', [attempt.run_id]), ('export', [attempt.run_id]),
                               ('attempt', [attempt.pk]), ('point_detail', [attempt.pk, point.pk])]:
                with self.subTest(page=name):
                    self.assertEqual(self.client.get(self.url(name, *args)).status_code, 200)
            check.assert_not_called()

    def test_student_page_retries_one_answer_without_wait_and_shows_fresh_result(self):
        from assignments.checks import formulas
        position = Position.objects.create(work=self.work, order=2)
        other_task = Task.objects.create(text_latex='Три плюс один?', answer_mask='{{number: 4}}')
        Candidate.objects.create(position=position, task=other_task)
        attempt = self.begin()
        for point in attempt.variant.points.all():
            Answer.objects.create(attempt=attempt, point=point, values={'s1': '4'},
                                  version=1, auto_check='')
        def checked(spec, values):
            self.assertEqual(formulas._worker_wait.get(), 0)
            return 'ok', {'s1': 'ok'}
        self.client.force_login(self.student)
        with patch.object(svc, 'check_mask', side_effect=checked) as check:
            response = self.client.get(self.url('attempt', attempt.pk))
            check.assert_called_once()
        self.assertEqual(response.status_code, 200)
        self.assertCountEqual([row['answer'].auto_check for row in response.context['rows']], ['ok', ''])
        self.assertIsNone(formulas._worker_wait.get())

    def test_manage_recipients_in_shared_dialog(self):
        a = self.issued()
        extra = get_user_model().objects.create_user('extra_recipient')
        Profile.objects.create(user=extra, teacher=self.teacher, role='student')
        self.client.force_login(self.teacher)
        page = self.client.get(reverse('assignment_list'))
        self.assertContains(page, f'data-dialog-open="assignment-group-independent-{a.run_id}"')
        self.assertContains(page, self.url('recipients', a.run_id))
        response = self.client.post(self.url('recipients', a.run_id), {'students': [extra.pk], f'variant_{extra.pk}': a.variant_id})
        self.assertEqual(response.status_code, 302)
        self.assertFalse(Attempt.objects.filter(pk=a.pk).exists())
        new = Attempt.objects.get(run=a.run)
        self.assertEqual(new.student_id, extra.pk)
        self.assertEqual(new.variant_id, a.variant_id)
        a.run.refresh_from_db()
        self.assertEqual(a.run.plan['students'], {str(extra.pk): a.variant.number})
        self.client.force_login(self.student)
        self.assertEqual(self.client.post(self.url('recipients', a.run_id), {}).status_code, 403)

    def test_recipient_changes_preserve_started_attempts_and_check_access(self):
        a = self.begin()
        self.client.force_login(self.teacher)
        url = self.url('recipients', a.run_id)
        self.assertEqual(self.client.post(url, {'students': [self.other.pk], f'variant_{self.other.pk}': a.variant_id}).status_code, 403)
        extra = get_user_model().objects.create_user('started_extra')
        Profile.objects.create(user=extra, teacher=self.teacher, role='student')
        self.assertEqual(self.client.post(url, {'students': [extra.pk], f'variant_{extra.pk}': a.variant_id}).status_code, 409)
        foreign = self.issued().variant
        self.assertEqual(self.client.post(url, {'students': [self.student.pk], f'variant_{self.student.pk}': foreign.pk}).status_code, 400)
        a.refresh_from_db()
        self.assertIsNotNone(a.started_at)
        self.assertEqual(a.run.attempts.count(), 1)

    def test_closed_run_and_repeat_cannot_gain_recipients(self):
        extra = get_user_model().objects.create_user('late_recipient')
        Profile.objects.create(user=extra, teacher=self.teacher, role='student')
        a = self.issued(deadline=timezone.now() + timedelta(minutes=1))
        a.run.deadline = timezone.now() - timedelta(minutes=1)
        a.run.save(update_fields=['deadline'])
        with self.assertRaises(svc.WorkError):
            svc.update_recipients(a.run_id, {self.student.pk, extra.pk}, {self.student.pk: a.variant_id, extra.pk: a.variant_id})
        self.assertEqual(a.run.attempts.count(), 1)
        b = self.begin()
        svc.submit(b.pk)
        b.refresh_from_db()
        run = svc.prepare(b.run.work, self.teacher, [self.student.pk], parent=b, kind='correction')
        svc.issue(run.pk)
        variant = run.variants.get()
        with self.assertRaises(svc.WorkError):
            svc.update_recipients(run.pk, {extra.pk}, {extra.pk: variant.pk})
        self.assertEqual(run.attempts.get().student_id, self.student.pk)

    def test_cumulative_label_requires_a_correction(self):
        a = self.begin()
        svc.submit(a.pk)
        self.client.force_login(self.teacher)
        url = self.url('progress', a.run_id)
        self.assertNotContains(self.client.get(url), 'С учётом работы над ошибками:')
        a.refresh_from_db()
        correction = svc.prepare(a.run.work, self.teacher, [self.student.pk], parent=a, kind='correction')
        svc.issue(correction.pk)
        svc.start(correction.attempts.get().pk)
        self.assertContains(self.client.get(url), 'С учётом работы над ошибками:')
        self.assertContains(self.client.get(url), 'Для каждого исходного пункта берётся лучший зачёт')

    def test_replacement_preview_keeps_old_plan_but_blocks_old_issue(self):
        first = svc.prepare(self.work, self.teacher, [self.student.pk])
        old_plan = first.plan
        with self.assertRaises(svc.WorkError):
            svc.prepare(self.work, self.teacher, [], previous=first)
        first.refresh_from_db()
        self.assertIsNone(first.superseded_at)
        second = svc.prepare(self.work, self.teacher, [self.student.pk], duration=20, previous=first)
        first.refresh_from_db()
        self.assertIsNotNone(first.superseded_at)
        self.assertEqual(first.plan, old_plan)
        with self.assertRaises(svc.WorkError):
            svc.issue(first.pk)
        with self.assertRaises(svc.WorkError):
            svc.prepare(self.work, self.teacher, [self.student.pk], previous=first)
        svc.issue(second.pk)
        with self.assertRaises(svc.WorkError):
            svc.prepare(self.work, self.teacher, [self.student.pk], previous=second)
        second.refresh_from_db()
        self.assertIsNone(second.superseded_at)
        self.assertEqual(second.attempts.count(), 1)
        self.assertEqual(svc.issue(second.pk).pk, second.pk)

    def test_resume_deadline_native_format_and_replacement(self):
        deadline = (timezone.now() + timedelta(days=1)).replace(second=0, microsecond=0)
        first = svc.prepare(self.work, self.teacher, [self.student.pk], deadline=deadline)
        self.client.force_login(self.teacher)
        url = reverse('assignment_create') + f'?mode=independent&work={self.work.pk}&resume={first.pk}'
        response = self.client.get(url)
        value = timezone.localtime(deadline).strftime('%Y-%m-%dT%H:%M')
        self.assertContains(response, f'value="{value}"')
        response = self.client.post(url, {'duration': 40, 'deadline': value, 'students': [self.student.pk], 'allow_photos': 'on', 'release_mode': 'manual'})
        self.assertEqual(response.status_code, 302)
        first.refresh_from_db()
        self.assertIsNotNone(first.superseded_at)
        self.assertEqual(Run.objects.latest('pk').deadline, deadline)

    def test_saved_draft_is_available_from_independent_catalog(self):
        self.client.force_login(self.teacher)
        url = reverse('assignment_create') + (
            f'?mode=independent&work={self.work.pk}')
        response = self.client.post(url, {
            'mode': 'independent',
            'work': self.work.pk,
            'students': [self.student.pk],
            'duration': 35,
            'max_variants': 1,
            'swap_numbers': 'on',
            'release_mode': 'manual',
            'action': 'draft',
        })
        self.assertRedirects(response, self.url('catalog'))
        draft = Run.objects.get()
        self.assertIsNone(draft.issued_at)
        self.assertIsNone(draft.superseded_at)
        self.assertEqual(draft.duration, 35)
        self.assertEqual(draft.max_variants, 1)
        self.assertEqual((draft.swap_numbers, draft.restrict_clipboard), (True, False))
        self.assertFalse(draft.allow_photos)
        self.assertEqual(draft.plan['students'], {str(self.student.pk): 1})
        catalog = self.client.get(self.url('catalog'))
        resume_url = (reverse('assignment_create')
                      + f'?mode=independent&amp;work={self.work.pk}'
                      + f'&amp;resume={draft.pk}')
        self.assertContains(catalog, 'Черновики выдачи')
        self.assertContains(catalog, resume_url)
        resumed = self.client.get(
            reverse('assignment_create'), {
                'mode': 'independent', 'work': self.work.pk,
                'resume': draft.pk,
            })
        self.assertEqual(resumed.context['draft_saved'], True)
        self.assertContains(resumed, 'value="35"')
        self.assertContains(resumed, 'id="id_swap_numbers" checked')
        self.assertNotContains(resumed, 'id="id_restrict_clipboard" checked')
        self.assertNotContains(resumed, 'id_allow_photos" checked')

    def test_catalog_redirect_keeps_student_kind(self):
        self.client.force_login(self.student)
        self.assertRedirects(self.client.get(self.url('catalog')), reverse('dashboard') + '?kind=independent')

    def test_global_stats_reuse_loaded_runs_despite_search(self):
        from .integration import assignment_groups, assignment_runs
        self.issued()
        runs = assignment_runs(self.teacher, {})
        with self.assertNumQueries(0):
            self.assertEqual(assignment_groups(self.teacher, {'q': 'Не совпадает'}, runs=runs), [])
            self.assertEqual(assignment_groups(self.teacher, {}, runs=runs)[0]['students_total'], 1)
        self.client.force_login(self.teacher)
        response = self.client.get(reverse('assignment_list'), {'q': 'Не совпадает', 'kind': 'independent'})
        self.assertEqual(response.context['stats']['total'], 1)
        self.assertEqual(len(response.context['page'].object_list), 0)

    def test_common_review_dialog_and_independent_comment_save(self):
        a = self.begin()
        point = a.variant.points.get()
        svc.submit(a.pk)
        self.client.force_login(self.teacher)
        url = self.url('review', a.pk, point.pk)
        response = self.client.post(url, {'grade': '+'}, HTTP_ACCEPT='application/json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['result']['done'], 1)
        self.assertEqual(response.json()['cumulative']['done'], 1)
        self.assertIn('Полностью правильно', response.json()['feedback'])
        self.client.post(url, {f'p{point.pk}_comment': 'Проверь вычисления'}, HTTP_ACCEPT='application/json')
        answer = Answer.objects.get(attempt=a, point=point)
        self.assertEqual(answer.grade, '+')
        self.assertContains(self.client.get(self.url('progress', a.run_id)), 'progress-mark progress-ok')
        self.client.post(url, {'grade': ''}, HTTP_ACCEPT='application/json')
        answer.refresh_from_db()
        self.assertEqual(answer.comment, 'Проверь вычисления')
        detail = self.client.get(self.url('point_detail', a.pk, point.pk))
        self.assertContains(detail, 'data-grade-url')
        self.assertContains(detail, 'data-preserve-input')
        progress = self.client.get(self.url('progress', a.run_id))
        self.assertTemplateUsed(progress, 'assignments/_progress_table.html')
        self.assertContains(progress, self.url('point_detail', a.pk, point.pk))
        self.client.force_login(self.student)
        self.assertEqual(self.client.get(self.url('point_detail', a.pk, point.pk)).status_code, 403)

    def test_attempt_points_in_one_sheet(self):
        """Номера работы — одним листом (.iw-sheet), не карточкой на номер;
        у сданной работы пояснение — строкой листа."""
        Position.objects.create(work=self.work, order=2, text='Пояснение к блоку')
        task = Task.objects.create(text_latex='Три плюс три?', answer_mask='{{number: 6}}')
        Candidate.objects.create(position=Position.objects.create(work=self.work, order=3), task=task)
        a = self.begin()
        svc.submit(a.pk)
        self.client.force_login(self.student)
        response = self.client.get(self.url('attempt', a.pk))
        self.assertContains(response, 'class="card ws-view iw-sheet"', count=1)
        self.assertContains(response, 'class="iw-point"', count=2)
        self.assertContains(response, '<div class="math-content iw-note">Пояснение к блоку</div>', html=False)
        self.assertNotContains(response, 'iw-section')
        self.assertNotContains(response, 'data-iw-focus')

    def test_attempt_one_task_screen_for_student(self):
        """Работа ученика (макет A2): своя шапка с «Вариант N» и таймером,
        лента номеров с состоянием ответа и «Сдать работу» в той же строке,
        счёт ответов под названием, «Задание N», под карточкой — только
        следующее задание (22.10.2026); у учителя страница прежняя."""
        Position.objects.create(work=self.work, order=2, text='Пояснение к блоку')
        task = Task.objects.create(text_latex='Три плюс три?', answer_mask='{{number: 6}}')
        Candidate.objects.create(position=Position.objects.create(work=self.work, order=3), task=task)
        a = self.begin()
        first, second = a.variant.points.order_by('position')
        svc.save_answer(a.pk, first.pk, 0, {'s1': '4'})
        self.client.force_login(self.student)
        response = self.client.get(self.url('attempt', a.pk))
        self.assertContains(response, 'data-iw-focus')
        self.assertContains(
            response, f'Вариант {a.variant.number}, <span data-iw-answered role="status">1 ответ</span> из 2</span>',
            html=False)
        self.assertContains(response, 'data-iw-timer')
        self.assertContains(response, f'data-iw-strip="{first.pk}" data-state="saved"')
        self.assertContains(response, f'data-iw-strip="{second.pk}" data-state="empty"')
        self.assertContains(response, 'Задание 1</h2>')
        self.assertContains(response, '<div class="math-content iw-point-kicker">Пояснение к блоку</div>', html=False)
        self.assertNotContains(response, 'Есть ')
        self.assertNotContains(response, 'Ваш ответ')
        self.assertContains(response, 'class="iw-head-submit"', count=1)
        self.assertNotContains(response, 'iw-submit-bar')
        # Под карточкой — кнопка следующего номера; у последнего её нет
        self.assertContains(response, f'data-iw-nav="{first.pk}"', count=1)
        self.assertNotContains(response, f'data-iw-nav="{second.pk}"')
        self.assertNotContains(response, 'sidebar-section')
        self.assertNotContains(response, 'page-crumbs')
        self.client.force_login(self.teacher)
        response = self.client.get(self.url('attempt', a.pk))
        self.assertContains(response, 'page-crumbs')
        self.assertNotContains(response, 'data-iw-focus')

    def test_unified_student_dashboard_disclosure_and_closed_start(self):
        a = self.issued(deadline=timezone.now() + timedelta(minutes=1))
        self.client.force_login(self.student)
        response = self.client.get(reverse('dashboard'))
        self.assertContains(response, self.url('attempt', a.pk))
        self.assertNotContains(response, 'SECRET_ANSWER')
        with patch('independentworks.integration.timezone.now', return_value=timezone.now() + timedelta(minutes=2)):
            response = self.client.get(reverse('dashboard'))
            self.assertEqual(len(response.context['closed_unstarted']), 1)
            self.assertEqual(
                response.context['closed_unstarted'][0].get_status_display,
                'Не начато')
            self.assertIsNone(response.context['closed_unstarted'][0].result)
            self.assertEqual(response.context['todo'], [])
        a = svc.start(a.pk)
        point = a.variant.points.get()
        svc.save_answer(a.pk, point.pk, 0, {'s1': '4'})
        svc.submit(a.pk)
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(len(response.context['submitted']), 1)
        self.assertIsNone(response.context['submitted'][0].result)
        a.run.release_mode = 'auto'; a.run.release_flags = ['checks']
        a.run.save(update_fields=['release_mode', 'release_flags'])
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.context['reviewed'][0].result['percent'], 100)

    def test_mixed_catalog_pagination_and_material_types(self):
        from worksheets.models import Worksheet
        from assignments.models import Assignment
        sheet = Worksheet.objects.create(owner=self.teacher, title='Обычный листок')
        Assignment.objects.create(worksheet=sheet, student=self.student)
        for _ in range(31):
            self.issued()
        self.client.force_login(self.teacher)
        response = self.client.get(reverse('assignment_list'))
        self.assertEqual(response.context['page'].paginator.count, 32)
        first = {g['key'] for g in response.context['worksheet_groups']}
        response = self.client.get(reverse('assignment_list'), {'page': 2})
        second = {g['key'] for g in response.context['worksheet_groups']}
        self.assertFalse(first & second)
        self.assertEqual(len(first | second), 32)
        response = self.client.get(reverse('assignment_list'), {'kind': 'ordinary'})
        self.assertEqual(response.context['page'].paginator.count, 1)
        response = self.client.get(reverse('assignment_list'), {'kind': 'independent', 'q': 'Уравнения'})
        self.assertEqual(response.context['page'].paginator.count, 31)
        response = self.client.get(reverse('worksheet_list'))
        self.assertNotContains(response, self.url('edit', self.work.pk))
        self.assertEqual([w.pk for w in response.context['worksheets']], [sheet.pk])
        self.assertNotContains(response, 'Вид материала')
        self.assertNotContains(response, 'Создать самостоятельную')
        self.assertContains(self.client.get(self.url('catalog')), '<h1>Самостоятельные</h1>', html=True)
        self.assertRedirects(self.client.get(self.url('prepare', self.work.pk)), reverse('assignment_create') + f'?mode=independent&work={self.work.pk}')

    def test_idempotent_issue_start(self):
        a = self.issued()
        svc.issue(a.run_id)
        self.assertEqual(Attempt.objects.count(), 1)
        first, second = svc.start(a.pk), svc.start(a.pk)
        self.assertEqual(first.ends_at, second.ends_at)
        self.assertEqual(first.started_at, second.started_at)

    def test_snapshot_and_stale_preview(self):
        run = svc.prepare(self.work, self.teacher, [self.student.pk])
        self.task.text_latex = 'Изменено'
        self.task.save()
        with self.assertRaises(svc.WorkError):
            svc.issue(run.pk)
        a = self.issued()
        self.task.text_latex = 'Снова изменено'
        self.task.save()
        self.assertEqual(a.variant.points.get().snapshot['text'], 'Изменено')

    def test_deadline_closed_browser_and_not_started(self):
        now = timezone.now()
        a = self.issued(deadline=now + timedelta(minutes=3), duration=40)
        begun = svc.start(a.pk)
        self.assertEqual(begun.ends_at, a.run.deadline)
        with patch('independentworks.services.timezone.now', return_value=now + timedelta(minutes=4)):
            svc.finish_expired(begun)
            self.assertEqual(begun.submitted_at, a.run.deadline)
            with self.assertRaises(svc.WorkError):
                svc.save_answer(a.pk, a.variant.points.get().pk, 0, {'s1': '4'})
        another = self.issued(deadline=now + timedelta(minutes=2))
        with patch('independentworks.services.timezone.now', return_value=now + timedelta(minutes=4)):
            with self.assertRaises(svc.WorkError):
                svc.start(another.pk)
            self.assertIsNone(svc.result(another)['percent'])

    def test_versions_clear_and_no_check_in_save_response(self):
        a = self.begin()
        point = a.variant.points.get()
        self.client.force_login(self.student)
        response = self.client.post(self.url('save', a.pk, point.pk), {'version': 0, f'p{point.pk}_s1': '4'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(set(response.json()), {'version', 'saved_at', 'state', 'attempt_state'})
        self.assertEqual(Answer.objects.get().auto_check, 'ok')
        response = self.client.post(self.url('save', a.pk, point.pk), {'version': 0, f'p{point.pk}_s1': '5'}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code, 409)
        self.assertEqual(Answer.objects.get().values['s1'], '4')
        svc.save_answer(a.pk, point.pk, 1, {'s1': ''})
        self.assertEqual(Answer.objects.get().values, {'s1': ''})
        self.assertEqual(Answer.objects.get().auto_check, 'wrong')
        self.assertFalse(Answer.objects.get().pending)

    def test_disclosure_before_start_submit_and_manual_flags(self):
        a = self.issued(release_mode='auto', release_flags=['checks', 'answer', 'solution'])
        self.client.force_login(self.student)
        self.assertNotContains(self.client.get(self.url('attempt', a.pk)), self.task.text_latex)
        svc.start(a.pk)
        response = self.client.get(self.url('attempt', a.pk))
        self.assertContains(response, self.task.text_latex)
        self.assertNotContains(response, 'SECRET_ANSWER')
        self.assertNotContains(response, 'SECRET_SOLUTION')
        svc.submit(a.pk)
        self.assertContains(self.client.get(self.url('attempt', a.pk)), 'SECRET_ANSWER')
        a.run.release_flags = ['answer']
        a.run.save(update_fields=['release_flags'])
        response = self.client.get(self.url('attempt', a.pk))
        self.assertContains(response, 'SECRET_ANSWER')
        self.assertNotContains(response, 'SECRET_SOLUTION')
        self.assertNotContains(response, 'Зачтено')

    def test_foreign_attempt_point_and_teacher_routes(self):
        a = self.begin()
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(self.url('attempt', a.pk)).status_code, 404)
        self.assertEqual(self.client.post(self.url('start', a.pk)).status_code, 404)
        self.assertEqual(self.client.post(self.url('save', a.pk, a.variant.points.get().pk)).status_code, 404)
        self.client.force_login(self.student)
        b = self.issued()
        self.assertEqual(self.client.post(self.url('save', a.pk, b.variant.points.get().pk)).status_code, 404)
        for name in ['progress', 'export', 'preview']:
            self.assertEqual(self.client.get(self.url(name, a.run_id)).status_code, 403)

    def two_parts(self):
        return [TaskSubitem.objects.create(task=self.task, order=i, text_latex=str(i), answer_mask='{{number: ' + str(i) + '}}') for i in (2, 3)]

    def test_same_task_separate_positions_and_partial_credit(self):
        first, second = self.two_parts()
        self.candidate.subitem_ids = [first.pk]
        self.candidate.save()
        p = Position.objects.create(work=self.work, order=2)
        Candidate.objects.create(position=p, task=self.task, subitem_ids=[second.pk])
        a = self.begin()
        points = list(a.variant.points.all())
        self.assertEqual(len(points), 2)
        svc.save_answer(a.pk, points[0].pk, 0, {'s1': '2'})
        svc.save_answer(a.pk, points[1].pk, 0, {'s1': '3'})
        self.assertEqual(svc.result(a)['done'], 2)
        Answer.objects.filter(point=points[1]).update(grade='+-')
        self.assertEqual(svc.result(a)['done'], 1.75)
        self.assertEqual(svc.result(a)['percent'], 88)

    def test_duplicate_point_rejected(self):
        Candidate.objects.create(position=Position.objects.create(work=self.work, order=2), task=self.task)
        with self.assertRaises(svc.WorkError):
            self.issued()

    def test_mismatch_confirmation(self):
        self.two_parts()
        Candidate.objects.create(position=self.position, task=Task.objects.create(text_latex='Другой кандидат'))
        with self.assertRaises(svc.WorkError):
            self.issued()
        self.position.mismatch_confirmed = True
        self.position.confirmed_composition = svc.composition(self.position)
        self.position.save()
        self.issued()

    def test_correction_snapshot_only_error_history(self):
        self.two_parts()
        a = self.begin()
        p, q = list(a.variant.points.all())
        svc.save_answer(a.pk, p.pk, 0, {'s1': '2'})
        svc.save_answer(a.pk, q.pk, 0, {'s1': '4'})
        svc.submit(a.pk)
        a.refresh_from_db()
        self.task.text_latex = 'Условие изменено'
        self.task.save()
        run = svc.prepare(self.work, self.teacher, [self.student.pk], parent=a, kind='correction')
        svc.issue(run.pk)
        b = svc.start(run.attempts.get().pk)
        point = b.variant.points.get()
        self.assertEqual(point.source_key, q.source_key)
        self.assertEqual(point.snapshot, q.snapshot)
        svc.save_answer(b.pk, point.pk, 0, {'s1': '3'})
        self.assertEqual(svc.result(a)['done'], 1)
        self.assertEqual(svc.cumulative(a)['done'], 2)

    def test_retake_avoids_seen(self):
        Candidate.objects.create(position=self.position, task=Task.objects.create(text_latex='Аналог', answer_mask='{{number: 5}}'))
        a = self.begin()
        svc.submit(a.pk); a.refresh_from_db()
        run = svc.prepare(self.work, self.teacher, [self.student.pk], parent=a, kind='retake')
        self.assertNotEqual(run.plan['variants'][0][0]['source_key'], a.variant.points.get().source_key)

    def test_shortage_and_cap(self):
        run = svc.prepare(self.work, self.teacher, [self.student.pk, self.other.pk])
        with self.assertRaises(svc.WorkError):
            svc.issue(run.pk)
        svc.issue(run.pk, confirm=True)
        self.assertEqual(run.attempts.count(), 2)
        capped = svc.prepare(self.work, self.teacher, [self.student.pk, self.other.pk], max_variants=1)
        self.assertFalse(capped.plan['shortage'])
        svc.issue(capped.pk)
        self.assertEqual(capped.variants.count(), 1)

    def test_started_reassignment_forbidden(self):
        a = self.begin()
        self.client.force_login(self.teacher)
        self.assertEqual(self.client.post(self.url('reassign', a.pk), {'variant': a.variant_id}).status_code, 409)

    def test_assets_private_and_independent(self):
        stream = io.BytesIO()
        Image.new('RGB', (10, 10)).save(stream, 'PNG')
        image = TaskImage.objects.create(task=self.task, image=ContentFile(stream.getvalue(), name='figure.png'))
        self.task.solution_latex = r'SECRET_SOLUTION \taskpic{1}'
        self.task.save()
        self.client.force_login(self.teacher)
        preview_run = svc.prepare(self.work, self.teacher, [self.student.pk])
        preview = self.client.get(self.url('preview', preview_run.pk))
        self.assertContains(preview, self.url('preview_asset', preview_run.pk, 1, 1, 1))
        self.assertTemplateUsed(preview, 'problems/_task_images.html')
        a = self.begin(release_mode='auto', release_flags=['solution'])
        asset = Asset.objects.get()
        self.client.force_login(self.student)
        self.assertEqual(self.client.get(self.url('asset', a.pk, asset.pk)).status_code, 404)
        image.image.delete(); image.delete()
        svc.submit(a.pk)
        response = self.client.get(self.url('asset', a.pk, asset.pk))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(b''.join(response.streaming_content))

    def test_picture_from_hidden_answer_version_not_shown_as_statement(self):
        # Авторский ответ текстом, рисунок — только в редакторской версии:
        # это не «картинка без ссылок», ученик не видит её до раскрытия.
        stream = io.BytesIO()
        Image.new('RGB', (10, 10)).save(stream, 'PNG')
        TaskImage.objects.create(task=self.task, image=ContentFile(stream.getvalue(), name='answer.png'))
        TaskImage.objects.create(task=self.task, image=ContentFile(stream.getvalue(), name='solution.png'))
        self.task.author_answer_latex = '4'
        self.task.answer_latex = r'4 \taskpic{1}'
        self.task.author_solution_latex = 'Считаем.'
        self.task.solution_latex = r'Считаем. \taskpic{2}'
        self.task.save()
        a = self.begin(release_mode='auto', release_flags=['answer', 'solution'])
        assets = {asset.name: asset for asset in Asset.objects.all()}
        self.assertEqual(assets['answer.png'].audiences, ['answer'])
        self.assertEqual(assets['solution.png'].audiences, ['solution'])
        self.client.force_login(self.student)
        page = self.client.get(self.url('attempt', a.pk))
        for asset in assets.values():
            self.assertNotContains(page, self.url('asset', a.pk, asset.pk))
            self.assertEqual(self.client.get(self.url('asset', a.pk, asset.pk)).status_code, 404)
        svc.submit(a.pk)
        page = self.client.get(self.url('attempt', a.pk))
        for asset in assets.values():
            self.assertContains(page, self.url('asset', a.pk, asset.pk))

    def test_migration_hides_answer_pictures_of_issued_variants(self):
        import importlib
        from django.apps import apps
        migration = importlib.import_module('independentworks.migrations.0011_hide_answer_assets')
        stream = io.BytesIO()
        Image.new('RGB', (10, 10)).save(stream, 'PNG')
        for name in ('figure.png', 'answer.png', 'loose.png'):
            TaskImage.objects.create(task=self.task, image=ContentFile(stream.getvalue(), name=name))
        self.task.text_latex += r' \taskpic{1}'
        self.task.author_answer_latex = '4'
        self.task.answer_latex = r'4 \taskpic{2}'
        self.task.save()
        a = self.issued()
        pending = svc.prepare(self.work, self.teacher, [self.student.pk])
        # Так их размечал снимок до исправления.
        Asset.objects.update(audiences=['statement'])
        for row in pending.plan['variants'][0]:
            for source in row['snapshot']['assets']:
                source['audiences'] = ['statement']
        pending.save(update_fields=['plan'])
        migration.forwards(apps, None)
        self.assertEqual({x.name: x.audiences for x in Asset.objects.filter(point__variant=a.variant)},
                         {'figure.png': ['statement'], 'answer.png': ['answer'], 'loose.png': ['statement']})
        pending.refresh_from_db()
        self.assertEqual({s['name']: s['audiences'] for s in pending.plan['variants'][0][0]['snapshot']['assets']},
                         {'figure.png': ['statement'], 'answer.png': ['answer'], 'loose.png': ['statement']})

    def test_migration_skips_points_whose_pictures_moved_after_issue(self):
        # Номера \taskpic снимка относятся к порядку картинок на момент
        # выдачи: если его поменяли, рисунок условия не трогаем.
        import importlib
        from django.apps import apps
        migration = importlib.import_module('independentworks.migrations.0011_hide_answer_assets')
        stream = io.BytesIO()
        Image.new('RGB', (10, 10)).save(stream, 'PNG')
        figure = TaskImage.objects.create(task=self.task, order=1, image=ContentFile(stream.getvalue(), name='figure.png'))
        TaskImage.objects.create(task=self.task, order=2, image=ContentFile(stream.getvalue(), name='answer.png'))
        self.task.text_latex += r' \taskpic{1}'
        self.task.author_answer_latex = '4'
        self.task.answer_latex = r'4 \taskpic{2}'
        self.task.save()
        a = self.issued()
        Asset.objects.update(audiences=['statement'])
        figure.order = 3
        figure.save(update_fields=['order'])
        migration.forwards(apps, None)
        self.assertEqual({x.name: x.audiences for x in Asset.objects.filter(point__variant=a.variant)},
                         {'figure.png': ['statement'], 'answer.png': ['statement']})

    def test_export_plain_text(self):
        self.student.first_name = '=1+1'; self.student.save()
        a = self.begin()
        self.client.force_login(self.teacher)
        response = self.client.get(self.url('export', a.run_id))
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(b''.join(response.streaming_content)))
        self.assertEqual(book.sheetnames, ['Сводка', 'Пункты'])
        self.assertEqual(book['Сводка']['A2'].data_type, 's')
        self.assertEqual(book['Сводка']['A2'].value, '=1+1')

    def test_pages_and_etag(self):
        self.client.force_login(self.teacher)
        run = svc.prepare(self.work, self.teacher, [self.student.pk])
        for name, args in [('create', []), ('edit', [self.work.pk]), ('pick', [self.position.pk]), ('prepare', [self.work.pk]), ('preview', [run.pk])]:
            with self.subTest(name=name):
                self.assertEqual(self.client.get(self.url(name, *args), follow=True).status_code, 200)
        svc.issue(run.pk)
        self.assertEqual(self.client.get(self.url('progress', run.pk)).status_code, 200)
        self.assertEqual(self.client.get(self.url('attempt', run.attempts.get().pk)).status_code, 200)
        response = self.client.get(self.url('progress', run.pk) + '?fragment=1')
        self.assertEqual(self.client.get(self.url('progress', run.pk) + '?fragment=1', HTTP_IF_NONE_MATCH=response['ETag']).status_code, 304)

    def test_pending_publication(self):
        self.task.answer_mask = ''; self.task.save()
        a = self.begin(); svc.submit(a.pk)
        self.client.force_login(self.teacher)
        self.client.post(self.url('publish', a.run_id), {'flags': ['checks']})
        self.client.force_login(self.student)
        self.assertContains(self.client.get(self.url('attempt', a.pk)), 'Задача на проверке')

    def test_feature_flag(self):
        self.client.force_login(self.teacher)
        with override_settings(INDEPENDENT_WORKS_ENABLED=False):
            self.assertEqual(self.client.get(self.url('catalog')).status_code, 404)

    def test_families_idempotent_and_pairs_separate(self):
        source = Source.objects.create(title='Зив')
        ids = []
        for v in range(1, 5):
            task = Task.objects.create(text_latex=f'Вариант {v}')
            ids.append(task.pk)
            TaskSource.objects.create(task=task, source=source, locator=f'sr/01/v{v}/3')
        bank = Source.objects.create(title='ФИПИ')
        for n in range(2):
            task = Task.objects.create(text_latex=f'Банк {n}')
            TaskSource.objects.create(task=task, source=bank, locator=f'2.1.{n}', source_path=['№2', 'Прототип 2.1', 'Аналоги'])
        from .analogs import sync_families
        sync_families(); sync_families()
        self.assertEqual(Family.objects.count(), 3)
        self.assertEqual(set(Task.objects.get(pk=ids[0]).analog_families.get().tasks.values_list('pk', flat=True)), set(ids[:2]))
        self.assertFalse(Task.objects.get(pk=ids[0]).analog_families.filter(tasks=ids[2]).exists())


    def test_block_candidates_remain_coherent(self):
        from problems.models import TaskBlock
        self.candidate.delete()
        source = Source.objects.create(title='Общие условия')
        blocks = []
        for variant in range(2):
            block = TaskBlock.objects.create(source=source, key=str(variant), title='Блок', size=5)
            members = [Task.objects.create(block=block, block_position=i, text_latex=f'Общее условие {variant}. Вопрос {i}.', answer_mask='{{number: 4}}') for i in range(1, 6)]
            svc.add_candidate(self.position, members[0])
            blocks.append({t.pk for t in members})
        self.assertEqual(self.work.positions.count(), 5)
        variants, shortage, _ = svc.generate(svc.material(self.work), 2, 'block')
        self.assertFalse(shortage)
        for points in variants:
            ids = {int(p['source_key'][1:]) for p in points}
            self.assertIn(ids, blocks)
        self.assertNotEqual(variants[0], variants[1])

    def test_oge_prototype_sections_do_not_merge(self):
        from .analogs import sync_families
        source = Source.objects.create(title='ФИПИ ОГЭ')
        tasks = []
        for prototype in (1, 2):
            for index, section in enumerate(('Прототипы', 'Аналоги')):
                task = Task.objects.create(text_latex=f'{prototype}/{index}')
                TaskSource.objects.create(task=task, source=source, locator=f'6.{prototype}.{index}', source_path=['№6', section])
                tasks.append(task)
        manual = Family.objects.create(key='manual:test', title='Ручная связь')
        manual.tasks.add(tasks[0], tasks[2])
        sync_families()
        sync_families()
        self.assertEqual(Family.objects.filter(origin='source').count(), 2)
        self.assertEqual(Family.objects.get(key=f'bank:{source.pk}:6.1').tasks.count(), 2)
        self.assertEqual(manual.tasks.count(), 2)

    def test_unpublished_correction_does_not_reveal_cumulative(self):
        a = self.begin(release_mode='auto', release_flags=['checks'])
        svc.submit(a.pk)
        a.refresh_from_db()
        run = svc.prepare(self.work, self.teacher, [self.student.pk], parent=a, kind='correction')
        svc.issue(run.pk)
        child = svc.start(run.attempts.get().pk)
        svc.save_answer(child.pk, child.variant.points.get().pk, 0, {'s1': '4'})
        svc.submit(child.pk)
        self.assertEqual(svc.cumulative(a)['percent'], 100)
        self.assertEqual(svc.cumulative(a, student_view=True)['percent'], 0)
        child.run.release_mode = 'auto'; child.run.release_flags = ['checks']
        child.run.save(update_fields=['release_mode', 'release_flags'])
        self.assertEqual(svc.cumulative(a, student_view=True)['percent'], 100)

    def test_manual_analog_finishes_in_normal_editor(self):
        self.client.force_login(self.teacher)
        response = self.client.post(self.url('create_analog', self.position.pk, self.task.pk))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Task.objects.count(), 1)
        editor = self.client.get(response.url)
        self.assertContains(editor, 'Сохранить аналог')
        self.assertEqual(Task.objects.count(), 1)
        data = {'text_latex': 'Сохранённый аналог', 'columns': 1, 'stars': 0}
        for name in ('formset', 'subitem_formset', 'image_formset'):
            prefix = editor.context[name].prefix
            data.update({prefix + '-TOTAL_FORMS': 0, prefix + '-INITIAL_FORMS': 0})
        self.assertEqual(self.client.post(response.url, data).status_code, 302)
        task = Task.objects.exclude(pk=self.task.pk).get()
        self.assertIsNone(task.variant_of_id)
        self.assertEqual(self.position.candidates.count(), 2)
        self.assertEqual(set(task.analog_families.get().tasks.values_list('pk', flat=True)), {task.pk, self.task.pk})

    def test_analog_from_worksheet_local_variant(self):
        local = self.task.make_variant(self.teacher)
        Candidate.objects.create(position=self.position, task=local)
        self.client.force_login(self.teacher)
        response = self.client.post(self.url('create_analog', self.position.pk, local.pk))
        self.assertEqual(response.status_code, 302)
        another = self.task.make_variant(self.teacher)
        self.assertEqual(self.client.post(self.url('create_analog', self.position.pk, another.pk)).status_code, 404)

    def test_photo_ownership_and_deadline(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Photo
        a = self.begin()
        p = a.variant.points.get()
        data = io.BytesIO()
        Image.new('RGB', (8, 8)).save(data, format='PNG')
        self.client.force_login(self.student)
        response = self.client.post(self.url('upload_photo', a.pk, p.pk), {'photo': SimpleUploadedFile('test.png', data.getvalue(), content_type='image/png')})
        self.assertEqual(response.status_code, 302)
        photo = Photo.objects.get()
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(self.url('photo', photo.pk)).status_code, 404)
        self.client.force_login(self.student)
        with patch('independentworks.models.timezone.now', return_value=a.ends_at):
            self.assertEqual(self.client.post(self.url('delete_photo', photo.pk)).status_code, 409)
            response = self.client.post(self.url('upload_photo', a.pk, p.pk), {'photo': SimpleUploadedFile('test.png', data.getvalue(), content_type='image/png')})
            self.assertEqual(response.status_code, 409)
        self.assertEqual(Photo.objects.count(), 1)


    def test_pasted_photo_from_clipboard(self):
        """Фото из буфера обмена: имя paste-….png и PNG без расширения."""
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Photo
        a = self.begin()
        p = a.variant.points.get()
        data = io.BytesIO()
        Image.new('RGB', (8, 8)).save(data, format='PNG')
        self.client.force_login(self.student)
        for name in ('paste-20261014-153000.png', 'paste'):
            response = self.client.post(self.url('upload_photo', a.pk, p.pk), {'photo': SimpleUploadedFile(name, data.getvalue(), content_type='image/png')})
            self.assertEqual(response.status_code, 302)
        self.assertEqual(Photo.objects.count(), 2)
        broken = self.client.post(self.url('upload_photo', a.pk, p.pk), {'photo': SimpleUploadedFile('paste.png', b'garbage', content_type='image/png')})
        self.assertContains(broken, 'Не удалось прочитать фотографию.', status_code=broken.status_code)
        self.assertNotEqual(broken.status_code, 500)

    def test_photo_fetch_returns_thumb_and_redirect_keeps_point(self):
        """Панель ответа грузит фото fetch'ем (JSON для миниатюры); без скрипта — назад к тому же заданию."""
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Photo
        a = self.begin()
        p = a.variant.points.get()
        data = io.BytesIO()
        Image.new('RGB', (8, 8)).save(data, format='PNG')
        self.client.force_login(self.student)
        page = self.client.get(self.url('attempt', a.pk))
        self.assertContains(page, 'data-iw-photo-input')
        self.assertContains(page, f'id="iw-photo-{p.pk}"')
        self.assertNotContains(page, 'Ваш ответ')
        xhr = {'HTTP_X_REQUESTED_WITH': 'XMLHttpRequest'}
        response = self.client.post(self.url('upload_photo', a.pk, p.pk), {'photo': SimpleUploadedFile('a.png', data.getvalue(), content_type='image/png')}, **xhr)
        photo = Photo.objects.get()
        self.assertEqual(response.json(), {'ok': True, 'id': photo.pk, 'url': self.url('photo', photo.pk), 'delete_url': self.url('delete_photo', photo.pk)})
        self.assertEqual(self.client.post(self.url('delete_photo', photo.pk), **xhr).json(), {'ok': True})
        self.assertFalse(Photo.objects.exists())
        response = self.client.post(self.url('upload_photo', a.pk, p.pk), {'photo': SimpleUploadedFile('a.png', data.getvalue(), content_type='image/png')})
        self.assertRedirects(response, self.url('attempt', a.pk) + f'#point-{p.pk}', fetch_redirect_response=False)
        response = self.client.post(self.url('delete_photo', Photo.objects.get().pk))
        self.assertRedirects(response, self.url('attempt', a.pk) + f'#point-{p.pk}', fetch_redirect_response=False)

    def test_save_rechecks_time_after_validation(self):
        from assignments.checks import check_mask
        a = self.begin()
        point = a.variant.points.get()
        clock = [a.started_at]
        def validate(spec, values):
            checked = check_mask(spec, values)
            clock[0] = a.ends_at
            return checked
        with patch('independentworks.services.timezone.now', side_effect=lambda: clock[0]):
            with patch('independentworks.services.check_mask', side_effect=validate):
                with self.assertRaises(svc.WorkError):
                    svc.save_answer(a.pk, point.pk, 0, {'s1': '4'})
        self.assertFalse(a.answers.exists())

    def test_shared_masks_through_save_endpoint(self):
        from assignments.checks import parse_mask
        cases = [
            ('number', '4', '4'), ('digits', '234', '234'),
            ('decimal', '-0,5', '-0,5'), ('ratio', '1:2', '1:2'),
            ('proportion', '1:2', '1:2'), ('simplify', 'x+1', 'x+1'),
            ('factor', '(x-1)(x+1)', '(x-1)(x+1)'),
            ('equation', 'x=2', 'x=2'), ('set', '-3; 3', '-3; 3'),
            ('pair', '(3; -4)', '(3; -4)'), ('range', '(-2; 3]', '(-2; 3]'),
            ('choice', 'a|*b', 'b'), ('order', 'a|b|c', '0|1|2'),
        ]
        self.client.force_login(self.student)
        for kind, reference, value in cases:
            with self.subTest(kind=kind):
                self.task.answer_mask = '{{' + kind + ': ' + reference + '}}'
                self.task.save()
                a = self.begin()
                p = a.variant.points.get()
                data = {'version': 0, f'p{p.pk}_s1': value}
                if kind == 'order':
                    tokens = parse_mask(p.snapshot['mask']).slots[0].order_tokens(f'p{p.pk}')
                    data = {'version': 0, **{f'p{p.pk}_s1_{i}': tokens[i] for i in range(3)}}
                response = self.client.post(self.url('save', a.pk, p.pk), data)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(a.answers.get().auto_check, 'ok')
                self.assertNotIn('checks', response.json())


    def test_mysql_settings_do_not_import_working_secrets(self):
        import os
        import subprocess
        import sys
        from pathlib import Path
        env = dict(os.environ, IW_TEST_MYSQL_HOST='localhost', IW_TEST_MYSQL_USER='isolated',
                   IW_TEST_MYSQL_PASSWORD='', DJANGO_ENV='production', LATEX_API_KEY='sentinel',
                   GEMINI_API_KEY='sentinel', DJANGO_SECRET_KEY='sentinel')
        code = "import sys; import independentworks.mysql_settings as s; assert 'core.settings' not in sys.modules; assert s.LATEX_API_KEY == s.GEMINI_API_KEY == ''; assert s.SECRET_KEY != 'sentinel'"
        result = subprocess.run([sys.executable, '-c', code], env=env,
                                cwd=Path(__file__).resolve().parent.parent, capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr.decode())
        del env['IW_TEST_MYSQL_USER']
        result = subprocess.run([sys.executable, '-c', 'import independentworks.mysql_settings'],
                                env=env, cwd=Path(__file__).resolve().parent.parent, capture_output=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn(b'IW_TEST_MYSQL_USER', result.stderr)

    def test_prepare_rejects_invalid_parent_kind_combinations(self):
        parent = self.begin()
        svc.submit(parent.pk)
        parent.refresh_from_db()
        for kind, previous in [('correction', None), ('retake', None), ('initial', parent), ('bogus', None)]:
            with self.subTest(kind=kind, parent=previous is not None):
                with self.assertRaises(svc.WorkError):
                    svc.prepare(self.work, self.teacher, [self.student.pk], kind=kind, parent=previous)
        self.client.force_login(self.teacher)
        response = self.client.get(self.url('prepare', self.work.pk), {'parent': parent.pk}, follow=True)
        self.assertEqual(response.status_code, 400)

    def test_missing_snapshot_record_and_file_abort_issue(self):
        for missing in ('record', 'file'):
            with self.subTest(missing=missing):
                a = self.begin()
                asset = Asset.objects.create(point=a.variant.points.get(), name='test.png', audiences=['statement'])
                asset.file.save('test.png', ContentFile(b'snapshot'))
                svc.submit(a.pk)
                a.refresh_from_db()
                run = svc.prepare(self.work, self.teacher, [self.student.pk], kind='correction', parent=a)
                if missing == 'record':
                    asset.delete()
                else:
                    asset.file.delete(save=False)
                with self.assertRaises(svc.WorkError) as caught:
                    svc.issue(run.pk)
                self.assertEqual(caught.exception.status, 409)
                self.assertFalse(run.variants.exists())
                self.assertFalse(run.attempts.exists())

    def test_long_family_key_preserves_title(self):
        from .analogs import sync_families
        source = Source.objects.create(title='Длинный прототип')
        prototype = '1' * 130 + '.' + '2' * 130
        for i in range(2):
            task = Task.objects.create(text_latex=str(i))
            TaskSource.objects.create(task=task, source=source, locator=str(i), source_path=['Прототип ' + prototype])
        sync_families()
        family = Family.objects.get(key__startswith='source:')
        self.assertTrue(family.title.startswith('Длинный прототип · прототип'))
        self.assertEqual(family.tasks.count(), 2)
        sync_families()
        self.assertEqual(Family.objects.filter(key__startswith='source:').count(), 1)

    def test_second_teacher_cannot_access_work_or_mutate_results(self):
        teacher = get_user_model().objects.create_user('teacher2')
        Profile.objects.create(user=teacher, role='teacher')
        a = self.begin()
        p = a.variant.points.get()
        asset = Asset.objects.create(point=p, name='test.png', audiences=['statement'])
        asset.file.save('test.png', ContentFile(b'snapshot'))
        svc.submit(a.pk)
        self.client.force_login(teacher)
        paths = [('progress', a.run_id), ('export', a.run_id), ('preview', a.run_id),
                 ('attempt', a.pk), ('asset', a.pk, asset.pk), ('edit', self.work.pk),
                 ('pick', self.position.pk), ('prepare', self.work.pk)]
        for name, *args in paths:
            with self.subTest(route=name):
                self.assertEqual(self.client.get(self.url(name, *args)).status_code, 404)
        for name, args, data in [
            ('publish', [a.run_id], {'flags': ['checks']}),
            ('review', [a.pk, p.pk], {'grade': '+'}),
            ('repeat', [a.pk], {'kind': 'correction'}),
        ]:
            with self.subTest(route=name):
                self.assertEqual(self.client.post(self.url(name, *args), data).status_code, 404)
        a.refresh_from_db()
        self.assertIsNone(a.published_at)
        self.assertFalse(a.answers.exists())

    def test_after_submit_settings_replace_globally(self):
        extra = get_user_model().objects.create_user('global_release_student')
        Profile.objects.create(user=extra, teacher=self.teacher, role='student')
        run = svc.prepare(self.work, self.teacher, [self.student.pk, extra.pk],
                          release_mode='auto', release_flags=['checks'])
        svc.issue(run.pk, confirm=True)
        attempts = list(run.attempts.order_by('pk'))
        a = svc.start(attempts[0].pk)
        other_attempt = svc.start(attempts[1].pk)
        p = a.variant.points.get()
        svc.save_answer(a.pk, p.pk, 0, {'s1': '4'})
        svc.submit(a.pk)
        svc.submit(other_attempt.pk)
        Answer.objects.filter(attempt=a).update(comment='SECRET_COMMENT')
        self.client.force_login(self.teacher)
        self.assertEqual(self.client.post(self.url('publish', a.run_id), {'flags': ['comment']}).status_code, 302)
        a.run.refresh_from_db()
        a.refresh_from_db()
        self.assertEqual(a.run.release_flags, ['comment'])
        self.assertEqual(a.disclosure, {'comment'})
        other_attempt.refresh_from_db()
        self.assertEqual(other_attempt.disclosure, {'comment'})
        self.assertEqual(self.client.post(self.url('publish', a.run_id), {'flags': ['answer']}).status_code, 302)
        self.assertEqual(self.client.post(self.url('publish', a.run_id), {}).status_code, 302)
        a.run.refresh_from_db()
        a.refresh_from_db()
        self.assertEqual(a.run.release_mode, 'manual')
        self.assertEqual(a.disclosure, set())
        self.assertNotContains(self.client.get(self.url('progress', a.run_id)), 'name="attempts"')
        b = self.begin()
        svc.submit(b.pk)
        Answer.objects.create(attempt=b, point=b.variant.points.get(), comment='ONLY_COMMENT')
        self.client.post(self.url('publish', b.run_id), {'flags': ['comment']})
        self.client.force_login(self.student)
        page = self.client.get(self.url('attempt', b.pk))
        self.assertContains(page, 'ONLY_COMMENT')
        for hidden in ('SECRET_ANSWER', 'SECRET_SOLUTION', 'Зачтено'):
            self.assertNotContains(page, hidden)

    def test_auto_disclosure_each_category_without_manual_publication(self):
        self.client.force_login(self.student)
        for flag, marker in [('checks', 'Зачтено'), ('answer', 'SECRET_ANSWER'), ('solution', 'SECRET_SOLUTION'), ('comment', 'AUTO_COMMENT')]:
            with self.subTest(flag=flag):
                a = self.begin(release_mode='auto', release_flags=[flag])
                svc.save_answer(a.pk, a.variant.points.get().pk, 0, {'s1': '4'})
                Answer.objects.filter(attempt=a).update(comment='AUTO_COMMENT')
                before = self.client.get(self.url('attempt', a.pk))
                self.assertNotContains(before, marker)
                svc.submit(a.pk)
                a.refresh_from_db()
                self.assertIsNone(a.published_at)
                page = self.client.get(self.url('attempt', a.pk))
                self.assertContains(page, marker)
                for hidden in {'Зачтено', 'SECRET_ANSWER', 'SECRET_SOLUTION', 'AUTO_COMMENT'} - {marker}:
                    self.assertNotContains(page, hidden)

    def test_create_from_collection_uses_current_family_relation(self):
        from worksheets.models import Worksheet, WorksheetItem, Collection
        source = Source.objects.create(title='Источник для самостоятельной')
        sheet = Worksheet.objects.create(owner=self.teacher, title='Парная подборка')
        WorksheetItem.objects.create(worksheet=sheet, task=self.task, order=1)
        collection = Collection.objects.create(worksheet=sheet, source=source, is_published=True)
        analog = Task.objects.create(text_latex='Аналог')
        family = Family.objects.create(key='manual:collection', title='Пара')
        family.tasks.add(self.task, analog)
        self.assertEqual(Family.tasks.through._meta.db_table, 'independentworks_familymember')
        self.client.force_login(self.teacher)
        self.assertEqual(self.client.get(self.url('create')).status_code, 200)
        response = self.client.post(self.url('create'), {'title': 'Из подборки', 'collection': collection.pk})
        self.assertEqual(response.status_code, 302)
        work = Work.objects.get(title='Из подборки')
        self.assertEqual(set(work.positions.get().candidates.values_list('task_id', flat=True)), {self.task.pk})

    def test_preparation_error_keeps_form_and_values(self):
        self.client.force_login(self.teacher)
        response = self.client.post(reverse('assignment_create') + f'?mode=independent&work={self.work.pk}', {
            'duration': 15, 'max_variants': 2, 'release_toggle': '1', 'withhold_answers': 'on',
        })
        self.assertContains(response, 'Выберите учеников.')
        self.assertEqual(response.context['form']['duration'].value(), '15')
        self.assertEqual(response.context['form']['max_variants'].value(), '2')
        self.assertContains(response, 'Сформировать и просмотреть')

    def test_release_toggle_maps_to_manual_and_auto(self):
        from .forms import RunForm
        for checked, expected in [(True, 'manual'), (False, 'auto')]:
            data = {'duration': 15, 'release_toggle': '1'}
            if checked:
                data['withhold_answers'] = 'on'
            form = RunForm(data, teacher=self.teacher)
            self.assertTrue(form.is_valid(), form.errors)
            self.assertEqual(form.cleaned_data['release_mode'], expected)

    def test_common_condition_once_but_retained_in_detail(self):
        self.task.text_latex = 'Общее условие уникальное'
        self.task.save()
        for i in range(2):
            TaskSubitem.objects.create(task=self.task, order=i, text_latex=f'Пункт {i}', answer_mask='{{number: 1}}')
        run = svc.prepare(self.work, self.teacher, [self.student.pk])
        self.client.force_login(self.teacher)
        response = self.client.get(self.url('preview', run.pk))
        self.assertContains(response, 'Общее условие уникальное', count=1)
        self.assertContains(response, 'Ответ и решение видите только вы')  # макет E_Preview
        self.assertNotContains(response, 'data-preview-select="1"')
        svc.issue(run.pk)
        attempt = svc.start(run.attempts.get().pk)
        response = self.client.get(self.url('attempt', attempt.pk))
        self.assertContains(response, 'Общее условие уникальное', count=1)
        for point in attempt.variant.points.all():
            self.assertContains(self.client.get(self.url('point_detail', attempt.pk, point.pk)), 'Общее условие уникальное', count=1)

    def test_add_students_preserves_started_attempt(self):
        attempt = self.begin()
        extra = get_user_model().objects.create_user('additional_student')
        Profile.objects.create(user=extra, teacher=self.teacher, role='student')
        self.client.force_login(self.teacher)
        response = self.client.post(self.url('add_students', attempt.run_id), {'students': [extra.pk], 'variant': attempt.variant_id})
        self.assertRedirects(response, self.url('progress', attempt.run_id))
        attempt.refresh_from_db()
        self.assertIsNotNone(attempt.started_at)
        self.assertEqual(attempt.run.attempts.count(), 2)
        self.client.post(self.url('add_students', attempt.run_id), {'students': [extra.pk], 'variant': attempt.variant_id})
        self.assertEqual(attempt.run.attempts.count(), 2)
        response = self.client.post(self.url('add_students', attempt.run_id), {'students': [self.other.pk], 'variant': attempt.variant_id})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(attempt.run.attempts.count(), 2)

    def test_candidate_results_group_students_by_variant(self):
        extra = get_user_model().objects.create_user('variant_group_student')
        Profile.objects.create(user=extra, teacher=self.teacher, role='student')
        run = svc.prepare(self.work, self.teacher, [self.student.pk, extra.pk], max_variants=1)
        svc.issue(run.pk)
        attempts = list(run.attempts.order_by('pk'))
        for attempt, value in zip(attempts, ('4', '3')):
            attempt = svc.start(attempt.pk)
            point = attempt.variant.points.get()
            svc.save_answer(attempt.pk, point.pk, 0, {'s1': value})
        self.client.force_login(self.teacher)
        response = self.client.get(self.url('progress', run.pk))
        self.assertEqual(len(response.context['variant_stats']), 1)
        stat = response.context['variant_stats'][0]
        self.assertEqual((stat['number'], stat['count'], stat['started']), (1, 2, 2))
        self.assertEqual((stat['done'], stat['points'], stat['percent']), (1, 2, 50))
        self.assertContains(response, 'Учеников 2')  # без «·»: «Учеников 2, начали …»
        self.assertContains(response, '1 / 2, 50%')

    def test_self_check_persists_and_logs_every_checked_answer(self):
        attempt = self.begin(self_check=False)
        self.client.force_login(self.teacher)
        url = self.url('self_check_settings', attempt.run_id)
        self.assertRedirects(self.client.post(url, {'self_check': 'on'}), self.url('progress', attempt.run_id))
        attempt.run.refresh_from_db()
        self.assertTrue(attempt.run.self_check)
        page = self.client.get(self.url('progress', attempt.run_id))
        self.assertContains(page, 'name="self_check" checked')
        self.client.force_login(self.student)
        point = attempt.variant.points.get()
        check_url = self.url('check_answer', attempt.pk, point.pk)
        response = self.client.post(check_url, {
            'version': 0, f'p{point.pk}_s1': '3'}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.json()['checks'][f'p{point.pk}_s1'], 'wrong')
        response = self.client.post(check_url, {
            'version': 1, f'p{point.pk}_s1': '4'}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.json()['checks'][f'p{point.pk}_s1'], 'ok')
        history = AnswerCheckAttempt.objects.filter(answer__attempt=attempt)
        self.assertEqual(list(history.values_list('number', 'result')), [(1, 'wrong'), (2, 'ok')])
        self.assertEqual(history.first().snapshot['values'], {'s1': '3'})
        from workaudit.models import AnswerRevision
        self.assertEqual(AnswerRevision.objects.filter(answer__attempt=attempt).count(), 2)
        self.client.force_login(self.teacher)
        teacher_page = self.client.get(self.url('attempt', attempt.pk))
        self.assertContains(teacher_page, 'Самопроверки: 2')
        self.assertContains(teacher_page, 'Поле 1: 3')
        self.assertRedirects(self.client.post(url, {}), self.url('progress', attempt.run_id))
        attempt.run.refresh_from_db()
        self.assertFalse(attempt.run.self_check)
        self.client.force_login(self.student)
        self.assertEqual(self.client.post(check_url, {
            'version': 2, f'p{point.pk}_s1': '4'}).status_code, 403)

    def test_parts_save_returns_json_without_redirect(self):
        part = TaskSubitem.objects.create(task=self.task, text_latex='Подпункт')
        self.client.force_login(self.teacher)
        response = self.client.post(self.url('pick', self.position.pk), {
            'task': self.task.pk, 'candidate': self.candidate.pk,
            'parts_mode': 'selected', 'parts': [part.pk],
        }, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['message'], 'Подпункты сохранены')
        self.candidate.refresh_from_db()
        self.assertEqual(self.candidate.subitem_ids, [part.pk])

    def test_parity_yo_has_no_duplicate_choices(self):
        from assignments.checks import parse_mask, check_slot
        for reference in ('нечетная', 'нечётная'):
            spec = parse_mask('{{parity: ' + reference + '}}')
            slot = spec.slots[0]
            self.assertEqual(slot.options, ['чётная', 'нечётная', 'общего вида'])
            for answer in ('нечетная', 'нечётная'):
                self.assertEqual(check_slot(slot, answer), 'ok')


from concurrent.futures import ThreadPoolExecutor
from django.db import connections, close_old_connections
from django.test import TransactionTestCase, skipUnlessDBFeature


@skipUnlessDBFeature('has_select_for_update')
@override_settings(INDEPENDENT_WORKS_ENABLED=True)
class ConcurrentWorkTests(TransactionTestCase):
    def setUp(self):
        User = get_user_model()
        self.teacher = User.objects.create_user('load_teacher')
        work = Work.objects.create(owner=self.teacher, title='Нагрузка')
        for number in range(10):
            position = Position.objects.create(work=work, order=number + 1)
            task = Task.objects.create(text_latex=f'Пункт {number + 1}: два плюс два', answer_mask='{{number: 4}}')
            Candidate.objects.create(position=position, task=task)
        User.objects.bulk_create([User(username=f'load_{i}') for i in range(30)])
        self.students = list(User.objects.filter(username__regex=r'^load_[0-9]+$'))
        self.run = svc.prepare(work, self.teacher, [u.pk for u in self.students], max_variants=1)
        svc.issue(self.run.pk)
        self.ids = list(self.run.attempts.values_list('pk', flat=True))
        self.point_ids = list(self.run.variants.get().points.values_list('pk', flat=True))
        self.point_id = self.point_ids[0]

    def test_parallel_issue_and_preview_replacement(self):
        from threading import Barrier
        from .models import Run
        old = svc.prepare(self.run.work, self.teacher, [self.students[0].pk])
        barrier = Barrier(2)
        def operation(replace):
            close_old_connections()
            try:
                barrier.wait(timeout=10)
                if replace:
                    svc.prepare(self.run.work, self.teacher, [self.students[0].pk], previous=old)
                    return 'replaced'
                svc.issue(old.pk)
                return 'issued'
            except svc.WorkError:
                return 'conflict'
            finally:
                connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:
            outcomes = list(pool.map(operation, [True, False]))
        old.refresh_from_db()
        self.assertEqual(outcomes.count('conflict'), 1)
        self.assertNotEqual(bool(old.issued_at), bool(old.superseded_at))
        self.assertEqual(old.attempts.count(), 1 if old.issued_at else 0)
        self.assertEqual(Run.objects.filter(work=self.run.work).count(), 3 if old.superseded_at else 2)

    def test_30_students_save_with_monitoring(self):
        def student(attempt_id):
            close_old_connections()
            try:
                svc.start(attempt_id)
                import time
                for version in range(3):
                    for point_id in self.point_ids:
                        time.sleep(0.002 * (attempt_id % 3))
                        svc.save_answer(attempt_id, point_id, version, {'s1': str(version + 2)})
                svc.submit(attempt_id)
                return attempt_id
            finally:
                connections.close_all()
        def monitor():
            close_old_connections()
            try:
                from .views import progress_context
                from django.test.utils import CaptureQueriesContext
                from django.db import connection
                for _ in range(5):
                    with CaptureQueriesContext(connection) as queries:
                        context = progress_context(self.run)
                    self.assertEqual(len(context['people']), 30)
                    self.assertEqual(len(context['numbers']), 10)
                    self.assertLess(len(queries), 100)
            finally:
                connections.close_all()
        with ThreadPoolExecutor(max_workers=31) as pool:
            watching = pool.submit(monitor)
            results = list(pool.map(student, self.ids))
            watching.result()
        self.assertEqual(len(results), 30)
        self.assertEqual(Answer.objects.filter(attempt__run=self.run, version=3, auto_check='ok').count(), 300)
        self.assertEqual(self.run.attempts.filter(submitted_at__isnull=False).count(), 30)

    def test_parallel_start_and_version_conflict(self):
        attempt_id = self.ids[0]
        def begin(_):
            close_old_connections()
            try:
                a = svc.start(attempt_id)
                return a.started_at, a.ends_at
            finally:
                connections.close_all()
        with ThreadPoolExecutor(max_workers=8) as pool:
            starts = list(pool.map(begin, range(8)))
        self.assertEqual(len(set(starts)), 1)
        def save(_):
            close_old_connections()
            try:
                svc.save_answer(attempt_id, self.point_id, 0, {'s1': '4'})
                return 'saved'
            except svc.WorkError:
                return 'conflict'
            finally:
                connections.close_all()
        with ThreadPoolExecutor(max_workers=8) as pool:
            results = list(pool.map(save, range(8)))
        self.assertEqual(results.count('saved'), 1)
        self.assertEqual(results.count('conflict'), 7)
